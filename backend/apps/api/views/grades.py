from rest_framework.viewsets import ModelViewSet
from rest_framework import permissions, status, serializers, parsers
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated, IsAdminUser
from django.db import transaction
from drf_spectacular.utils import extend_schema

from apps.grades.models import AssessmentGrade
from apps.api.serializers.grades import (
    AssessmentGradeSerializer,
    AssessmentGradeDetailSerializer,
    BulkGradeCreateSerializer,
    BulkGradeUpdateSerializer,
    StudentGradeReportSerializer,
    BulkGradeImportSerializer,
)
from apps.api.permissions import IsDepartmentHead, IsInstructor, IsStudent
from apps.grades.calculators import AchievementCalculator
from django.shortcuts import get_object_or_404
from apps.users.models import User
from apps.courses.models import CourseInstance

class GradeViewSet(ModelViewSet):
    """
    ViewSet for managing student grades.
    """
    queryset = AssessmentGrade.objects.all()
    
    def get_serializer_class(self):
        if self.action == 'retrieve':
            return AssessmentGradeDetailSerializer
        return AssessmentGradeSerializer

    def get_permissions(self):
        """
        Create/Update: Instructor of the course only.
        Read: Student (own), Instructor (course), Dept Head (dept).
        """
        if self.action in ['create', 'update', 'partial_update', 'destroy', 'bulk_create', 'bulk_update']:
            # Custom permission logic in perform_create/update or via custom permission class
            # For now using IsInstructor as base, specific check in methods
            permission_classes = [IsInstructor | IsDepartmentHead | IsAdminUser]
        else:
            permission_classes = [IsAuthenticated]
        return [permission() for permission in permission_classes]

    def get_queryset(self):
        """
        Filter grades based on user role.
        """
        queryset = super().get_queryset()
        user = self.request.user
        
        if user.is_department_head():
            # Department heads only see grades for courses in their department
            queryset = queryset.filter(
                assessment__course_instance__course_template__department=user.department
            )
        elif user.is_instructor():
            # Instructors see grades for their courses
            queryset = queryset.filter(assessment__course_instance__instructor=user)
        elif user.is_student():
            # Students see their own grades
            queryset = queryset.filter(student=user)
            
        return queryset

    def perform_create(self, serializer):
        """
        Validate that the instructor teaches the course for the assessment.
        """
        assessment = serializer.validated_data.get('assessment')
        user = self.request.user
        
        if not user.is_staff and assessment.course_instance.instructor != user:
            raise serializers.ValidationError("You can only enter grades for your own courses.")
            
        serializer.save(entered_by=user)

    def perform_update(self, serializer):
        """
        Validate that the instructor teaches the course for the assessment.
        """
        assessment = serializer.instance.assessment
        user = self.request.user
        
        if not user.is_staff and assessment.course_instance.instructor != user:
            raise serializers.ValidationError("You can only update grades for your own courses.")
            
        serializer.save()

    @extend_schema(
        request=BulkGradeCreateSerializer,
        responses={201: AssessmentGradeSerializer(many=True)},
        summary="Bulk create grades",
        description="Create multiple grades atomically. If any grade fails validation, the entire operation rolls back.",
        tags=["grades"]
    )
    @action(detail=False, methods=['post'])
    def bulk_create(self, request):
        """
        Bulk create grades for an assessment.
        All grades are created atomically - if one fails, all fail.
        """
        serializer = BulkGradeCreateSerializer(data=request.data, context={'request': request})
        serializer.is_valid(raise_exception=True)
        
        with transaction.atomic():
            created = []
            for grade_data in serializer.validated_data['grades']:
                grade = AssessmentGrade.objects.create(
                    assessment=grade_data['assessment_id'],
                    student=grade_data['student_id'],
                    score=grade_data['score'],
                    entered_by=request.user
                )
                created.append(grade)
        
        return Response(
            AssessmentGradeSerializer(created, many=True).data,
            status=status.HTTP_201_CREATED
        )

    @extend_schema(
        request=BulkGradeUpdateSerializer,
        responses={200: AssessmentGradeSerializer(many=True)},
        summary="Bulk update/upsert grades",
        description="Update or create multiple grades atomically. Uses upsert behavior - creates if not exists, updates if exists.",
        tags=["grades"]
    )
    @action(detail=False, methods=['post'])
    def bulk_update(self, request):
        """
        Bulk update/upsert grades for an assessment.
        Uses upsert behavior: creates new grades or updates existing ones.
        All operations are atomic - if one fails, all fail.
        """
        serializer = BulkGradeUpdateSerializer(data=request.data, context={'request': request})
        serializer.is_valid(raise_exception=True)
        
        with transaction.atomic():
            updated = []
            for grade_data in serializer.validated_data['grades']:
                grade, _ = AssessmentGrade.objects.update_or_create(
                    assessment=grade_data['assessment_id'],
                    student=grade_data['student_id'],
                    defaults={
                        'score': grade_data['score'],
                        'entered_by': request.user
                    }
                )
                updated.append(grade)
        
        return Response(
            AssessmentGradeSerializer(updated, many=True).data,
            status=status.HTTP_200_OK
        )

    @extend_schema(
        responses={200: StudentGradeReportSerializer(many=True)},
        summary="Get student grade report (all courses)",
        description="Get comprehensive grade reports for a student across all courses.",
        tags=["grades"]
    )
    @action(detail=False, methods=['get'], url_path='student/(?P<student_id>[^/.]+)/report')
    def student_report(self, request, student_id=None):
        student = None
        try:
            student = User.objects.get(student_id=student_id)
        except User.DoesNotExist:
            pass

        # Permission Check
        has_permission = False
        if student:
            is_self = request.user == student
            is_dept_head = request.user.is_department_head() and request.user.department == student.department
            if request.user.is_staff or is_dept_head or is_self:
                has_permission = True

        # If not found, only admins get 404. Others get 403 to prevent enumeration.
        if not student:
            if request.user.is_staff or request.user.is_superuser:
                return Response({"detail": "Student not found."}, status=status.HTTP_404_NOT_FOUND)
            return Response({"detail": "Permission denied."}, status=status.HTTP_403_FORBIDDEN)

        if not has_permission:
             return Response({"detail": "Permission denied."}, status=status.HTTP_403_FORBIDDEN)

        courses = student.enrolled_courses.select_related('course_template').all()
        reports = []
        
        for course in courses:
            # Construct report for each course
            grades = AssessmentGrade.objects.filter(student=student, assessment__course_instance=course).select_related('assessment')
            grade_info = AchievementCalculator.calculate_final_course_grade(student, course)
            
            reports.append({
                "student_id": student.id,
                "student_name": f"{student.first_name} {student.last_name}",
                "student_email": student.email,
                "course_instance_id": course.id,
                "course_name": course.course_template.name,
                "course_code": course.course_template.get_full_code(),
                "semester": course.semester,
                "year": course.year,
                "grades": grades, # Serializer handles queryset
                "total_weighted_score": grade_info['total_score'],
                "assessment_count": course.assessments.count(),
                "graded_count": grades.count()
            })
            
        serializer = StudentGradeReportSerializer(reports, many=True)
        return Response(serializer.data)

    @extend_schema(
        responses={200: StudentGradeReportSerializer},
        summary="Get course grade report for student",
        description="Get detailed grade report for a student in a specific course.",
        tags=["grades"]
    )
    @action(detail=False, methods=['get'], url_path='course/(?P<course_id>[^/.]+)/student/(?P<student_id>[^/.]+)/report')
    def course_student_report(self, request, course_id=None, student_id=None):
        try:
            student = User.objects.get(student_id=student_id)
            course = CourseInstance.objects.get(pk=course_id)
        except (User.DoesNotExist, CourseInstance.DoesNotExist):
             return Response({"detail": "Student or Course not found."}, status=status.HTTP_404_NOT_FOUND)

        # Permission checks
        is_self = request.user == student
        is_instructor = course.instructor == request.user
        is_dept_head = request.user.is_department_head() and request.user.department == course.course_template.department
        
        if not (is_instructor or is_dept_head or request.user.is_staff or is_self):
             return Response({"detail": "Permission denied."}, status=status.HTTP_403_FORBIDDEN)
             
        grades = AssessmentGrade.objects.filter(student=student, assessment__course_instance=course).select_related('assessment')
        grade_info = AchievementCalculator.calculate_final_course_grade(student, course)
        
        report = {
            "student_id": student.id,
            "student_name": f"{student.first_name} {student.last_name}",
            "student_email": student.email,
            "course_instance_id": course.id,
            "course_name": course.course_template.name,
            "course_code": course.course_template.get_full_code(),
            "semester": course.semester,
            "year": course.year,
            "grades": grades,
            "total_weighted_score": grade_info['total_score'],
            "assessment_count": course.assessments.count(),
            "graded_count": grades.count()
        }
        
        serializer = StudentGradeReportSerializer(report)
        return Response(serializer.data)

    @extend_schema(
        request=BulkGradeImportSerializer,
        responses={200: {"type": "object", "properties": {"processed": {"type": "integer"}, "message": {"type": "string"}}}},
        summary="Import grades from CSV/Excel",
        description="Bulk import grades by uploading a file mapping student_id to score.",
        tags=["grades"]
    )
    @action(detail=False, methods=['post'], parser_classes=[parsers.MultiPartParser])
    def import_grades(self, request):
        serializer = BulkGradeImportSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        
        file = serializer.validated_data['file']
        assessment_id = serializer.validated_data['assessment_id']
        
        # Verify assessment exists and instructor owns it
        from apps.courses.models import Assessment
        try:
            assessment = Assessment.objects.select_related('course_instance').get(id=assessment_id)
        except Assessment.DoesNotExist:
            return Response({"detail": "Assessment not found."}, status=status.HTTP_404_NOT_FOUND)
            
        if not request.user.is_staff and assessment.course_instance.instructor != request.user:
            return Response({"detail": "You can only import grades for your own courses."}, status=status.HTTP_403_FORBIDDEN)
            
        import csv
        import io
        
        grades_to_process = [] # List of dicts: {'student_id': str, 'score': val}
        
        try:
            if file.name.lower().endswith('.csv'):
                decoded_file = file.read().decode('utf-8')
                io_string = io.StringIO(decoded_file)
                reader = csv.DictReader(io_string)
                # Verify headers: fuzzy match or exact? Exact for now.
                # strip spaces from headers
                reader.fieldnames = [name.strip() for name in reader.fieldnames]
                
                if 'student_id' not in reader.fieldnames or 'score' not in reader.fieldnames:
                     return Response({"detail": "CSV must have 'student_id' and 'score' columns."}, status=status.HTTP_400_BAD_REQUEST)
                for row in reader:
                    if row['student_id'] and row['score']:
                        grades_to_process.append(row)
            else:
                # Excel
                import openpyxl
                wb = openpyxl.load_workbook(file, data_only=True)
                ws = wb.active
                
                headers = []
                for cell in ws[1]:
                    if cell.value:
                        headers.append(str(cell.value).strip())
                    else:
                        headers.append('')
                
                # Basic header check
                try:
                    sid_idx = headers.index('student_id')
                    score_idx = headers.index('score')
                except ValueError:
                     return Response({"detail": "Excel must have 'student_id' and 'score' columns."}, status=status.HTTP_400_BAD_REQUEST)
                
                for row in ws.iter_rows(min_row=2, values_only=True):
                    # Check if row is not empty
                    if row[sid_idx] is not None and row[score_idx] is not None:
                        grades_to_process.append({
                            'student_id': str(row[sid_idx]),
                            'score': row[score_idx]
                        })
        except Exception as e:
             return Response({"detail": f"Error parsing file: {str(e)}"}, status=status.HTTP_400_BAD_REQUEST)
             
        # Process grades
        processed_count = 0
        errors = []
        
        from django.contrib.auth import get_user_model
        User = get_user_model()
        
        with transaction.atomic():
            for idx, item in enumerate(grades_to_process):
                sid = str(item.get('student_id')).strip()
                score_val = item.get('score')
                
                try:
                    # Validate score is number
                    score = float(score_val)
                except (ValueError, TypeError):
                    errors.append(f"Row {idx+1}: Score '{score_val}' is not a valid number.")
                    continue

                if score > assessment.max_score:
                    errors.append(f"Row {idx+1}: Score {score} exceeds maximum {assessment.max_score}.")
                    continue
                
                try:
                    student = User.objects.get(student_id=sid, role="STUDENT")
                except User.DoesNotExist:
                    errors.append(f"Row {idx+1}: Student ID {sid} not found.")
                    continue
                    
                if not assessment.course_instance.students.filter(id=student.id).exists():
                    errors.append(f"Row {idx+1}: Student {sid} is not enrolled in this course.")
                    continue
                    
                # Upsert grade
                AssessmentGrade.objects.update_or_create(
                    assessment=assessment,
                    student=student,
                    defaults={'score': score, 'entered_by': request.user}
                )
                processed_count += 1
            
            if errors:
                # Rollback on any error for safety?
                # User might prefer partial success if list is huge?
                # But errors list is returned.
                # Let's rollback to ensure data integrity and user fixing file.
                transaction.set_rollback(True)
                return Response({
                    "detail": "Import failed due to errors.", 
                    "errors": errors[:50] # Limit error list size
                }, status=status.HTTP_400_BAD_REQUEST)
            
        return Response({"processed": processed_count, "message": f"Successfully imported {processed_count} grades."})

